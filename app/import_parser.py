"""Bounded XLSX validation and per-order import results."""

from datetime import date, datetime
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from pydantic import ValidationError

from app.erp import CreatePurchaseOrderRequest

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_ROWS = 2000
MAX_ERRORS = 100


def parse_and_import(source: BytesIO, store) -> dict:
    try:
        with ZipFile(source) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise ValueError("Excel 解壓後超過 50 MB，請拆分檔案")
    except BadZipFile as exc:
        raise ValueError("檔案不是有效的 XLSX") from exc
    source.seek(0)
    workbook = load_workbook(source, read_only=True, data_only=True)
    report = {"imported": 0, "skipped": 0, "failed": 0, "errors": [], "rows": 0}
    groups = {}
    invalid = set()
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = [str(v).strip().lower() if v is not None else "" for v in next(rows, ())]
        required = {
            "po_id",
            "supplier_name",
            "expected_date",
            "material_id",
            "material_name",
            "ordered_quantity",
        }
        if required - set(headers):
            raise ValueError(f"Excel 缺少欄位：{', '.join(sorted(required - set(headers)))}")
        nonempty = [h for h in headers if h]
        if len(nonempty) != len(set(nonempty)):
            raise ValueError("Excel 欄位名稱不可重複")
        positions = {h: i for i, h in enumerate(headers)}
        for row_number, row in enumerate(rows, 2):
            if row_number > MAX_ROWS + 1:
                raise ValueError(f"Excel 最多 {MAX_ROWS} 列，請拆分檔案")
            if not any(v is not None and str(v).strip() for v in row):
                continue
            report["rows"] += 1
            po_id = ""
            try:

                def cell(field, row=row):
                    index = positions.get(field)
                    value = row[index] if index is not None and index < len(row) else None
                    if value is None or not str(value).strip():
                        if field == "unit":
                            return "pcs"
                        raise ValueError(f"{field} 不可為空")
                    return value

                po_id = str(cell("po_id")).strip()
                expected = cell("expected_date")
                if isinstance(expected, datetime):
                    expected = expected.date()
                expected = (
                    expected.isoformat() if isinstance(expected, date) else str(expected).strip()
                )
                supplier = str(cell("supplier_name")).strip()
                item = {
                    "material_id": str(cell("material_id")).strip(),
                    "material_name": str(cell("material_name")).strip(),
                    "ordered_quantity": cell("ordered_quantity"),
                    "unit": str(cell("unit")),
                }
                group = groups.setdefault(
                    po_id,
                    {
                        "row": row_number,
                        "payload": {
                            "po_id": po_id,
                            "supplier_name": supplier,
                            "expected_date": expected,
                            "items": [],
                        },
                    },
                )
                if (supplier, expected) != (
                    group["payload"]["supplier_name"],
                    group["payload"]["expected_date"],
                ):
                    raise ValueError("同一採購單的供應商或到貨日期與前列不一致")
                # Validate each row before collecting it; do not truncate fractional quantities.
                validated = CreatePurchaseOrderRequest.model_validate(
                    {**group["payload"], "items": [item]}
                )
                group["payload"]["items"].append(validated.items[0].model_dump())
                CreatePurchaseOrderRequest.model_validate(group["payload"])
            except (ValueError, ValidationError) as exc:
                invalid.add(po_id or f"row:{row_number}")
                if len(report["errors"]) < MAX_ERRORS:
                    report["errors"].append(
                        {"row": row_number, "po_id": po_id, "message": str(exc)[:500]}
                    )
        if not report["rows"]:
            raise ValueError("Excel 沒有資料列")
        # File-wide size/row checks finish before any business records are created.
        for po_id, group in groups.items():
            if po_id in invalid:
                continue
            request = CreatePurchaseOrderRequest.model_validate(group["payload"])
            try:
                store.create_purchase_order(request)
                report["imported"] += 1
            except ValueError as exc:
                existing = store.repository.get_purchase_order(po_id)
                same = existing and (
                    existing.supplier_name == request.supplier_name
                    and existing.expected_date == request.expected_date
                    and [i.model_dump(exclude={"received_quantity"}) for i in existing.items]
                    == [i.model_dump() for i in request.items]
                )
                if same:
                    report["skipped"] += 1
                else:
                    invalid.add(po_id)
                    if len(report["errors"]) < MAX_ERRORS:
                        report["errors"].append(
                            {"row": group["row"], "po_id": po_id, "message": str(exc)[:500]}
                        )
        report["failed"] = len(invalid)
        return report
    finally:
        workbook.close()
