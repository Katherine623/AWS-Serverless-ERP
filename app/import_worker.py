from __future__ import annotations

import json
import logging
from datetime import date, datetime
from urllib.parse import unquote_plus

import boto3

from app.erp import CreatePurchaseOrderRequest, ErpStore

logger = logging.getLogger(__name__)
store = ErpStore()
s3 = boto3.client("s3")


def _normalise_date(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _import_workbook(bucket: str, key: str) -> int:
    from openpyxl import load_workbook

    response = s3.get_object(Bucket=bucket, Key=key)
    workbook = load_workbook(response["Body"], read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip().lower() if value is not None else "" for value in next(rows)]
        required = {
            "po_id",
            "supplier_name",
            "expected_date",
            "material_id",
            "material_name",
            "ordered_quantity",
        }
        missing = required - set(headers)
        if missing:
            raise ValueError(f"Excel 欄位不足: {', '.join(sorted(missing))}")
        index = {header: position for position, header in enumerate(headers)}
        grouped: dict[str, dict[str, object]] = {}
        for row in rows:
            if not any(value is not None and str(value).strip() for value in row):
                continue
            po_id = str(row[index["po_id"]]).strip()
            order = grouped.setdefault(
                po_id,
                {
                    "po_id": po_id,
                    "supplier_name": str(row[index["supplier_name"]]).strip(),
                    "expected_date": _normalise_date(row[index["expected_date"]]),
                    "items": [],
                },
            )
            order["items"].append(
                {
                    "material_id": str(row[index["material_id"]]).strip(),
                    "material_name": str(row[index["material_name"]]).strip(),
                    "ordered_quantity": int(row[index["ordered_quantity"]]),
                    "unit": str(row[index.get("unit", -1)] or "pcs").strip()
                    if "unit" in index
                    else "pcs",
                }
            )
        imported = 0
        for payload in grouped.values():
            store.create_purchase_order(CreatePurchaseOrderRequest.model_validate(payload))
            imported += 1
        return imported
    finally:
        workbook.close()


def handler(event: dict, context: object) -> dict[str, list[dict[str, str]]]:
    del context
    failures: list[dict[str, str]] = []
    for record in event.get("Records", []):
        message_id = record.get("messageId", "unknown")
        try:
            message = json.loads(record["body"])
            for s3_record in message.get("Records", []):
                bucket = s3_record["s3"]["bucket"]["name"]
                key = unquote_plus(s3_record["s3"]["object"]["key"])
                imported = _import_workbook(bucket, key)
                logger.info("Imported %s purchase orders from s3://%s/%s", imported, bucket, key)
        except Exception:
            logger.exception("Excel import message %s failed", message_id)
            failures.append({"itemIdentifier": message_id})
    return {"batchItemFailures": failures}
